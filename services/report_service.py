"""Exportacao de relatorios em Excel (.xlsx), CSV e PDF."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from services.calculation_service import DayCalculation
from services.exceptions import ServiceError
from utils.dates import format_date_br
from utils.formatting import format_minutes_as_hours, format_time_or_placeholder
from utils.money import format_brl

WORK_REPORT_HEADERS = [
    "Data", "Dia da semana", "Entrada", "Saida almoco", "Retorno", "Saida",
    "Horas trabalhadas", "Horas previstas", "Saldo", "Horas extras", "Tipo de dia",
    "Observacoes", "Status",
]


def build_work_report_rows(
    days: list[DayCalculation], records_by_date: dict
) -> list[dict[str, Any]]:
    from utils.dates import weekday_label

    rows = []
    for day in days:
        record = records_by_date.get(day.work_date)
        rows.append(
            {
                "Data": format_date_br(day.work_date),
                "Dia da semana": weekday_label(day.work_date),
                "Entrada": format_time_or_placeholder(record.entry_time if record else None),
                "Saida almoco": format_time_or_placeholder(record.lunch_start if record else None),
                "Retorno": format_time_or_placeholder(record.lunch_end if record else None),
                "Saida": format_time_or_placeholder(record.exit_time if record else None),
                "Horas trabalhadas": format_minutes_as_hours(day.worked_minutes),
                "Horas previstas": format_minutes_as_hours(day.expected_minutes),
                "Saldo": format_minutes_as_hours(day.balance_minutes, show_sign=True),
                "Horas extras": format_minutes_as_hours(day.overtime_minutes),
                "Tipo de dia": day.day_type.label_pt,
                "Observacoes": (record.notes or "") if record else "",
                "Status": "Completo" if day.is_complete else "Incompleto",
            }
        )
    return rows


FINANCE_REPORT_HEADERS = ["Periodo", "Horas normais", "Horas extras", "Valor hora", "Valor normal", "Valor extra", "Total estimado"]


def build_finance_report_rows(entries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """`entries` e uma lista de dicts com chaves: label, regular_minutes,
    overtime_minutes, hourly_rate, regular_value, overtime_value, total_value."""
    rows = []
    for e in entries:
        rows.append(
            {
                "Periodo": e["label"],
                "Horas normais": format_minutes_as_hours(e["regular_minutes"]),
                "Horas extras": format_minutes_as_hours(e["overtime_minutes"]),
                "Valor hora": format_brl(e["hourly_rate"]),
                "Valor normal": format_brl(e["regular_value"]),
                "Valor extra": format_brl(e["overtime_value"]),
                "Total estimado": format_brl(e["total_value"]),
            }
        )
    return rows


class ReportService:
    @staticmethod
    def export_csv(path: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
        try:
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                writer = csv.DictWriter(f, fieldnames=headers, delimiter=";")
                writer.writeheader()
                writer.writerows(rows)
        except OSError as exc:
            raise ServiceError("Nao foi possivel salvar o arquivo CSV.", technical_detail=str(exc)) from exc

    @staticmethod
    def export_xlsx(path: str, rows: list[dict[str, Any]], headers: list[str], title: str = "Relatorio") -> None:
        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = title[:31]

            sheet.append(headers)
            header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill

            for row in rows:
                sheet.append([row.get(h, "") for h in headers])

            for idx, header in enumerate(headers, start=1):
                max_length = max(
                    [len(str(header))] + [len(str(row.get(header, ""))) for row in rows]
                )
                sheet.column_dimensions[get_column_letter(idx)].width = min(max_length + 4, 40)

            Path(path).parent.mkdir(parents=True, exist_ok=True)
            workbook.save(path)
        except OSError as exc:
            raise ServiceError("Nao foi possivel salvar o arquivo Excel.", technical_detail=str(exc)) from exc

    @staticmethod
    def export_pdf(path: str, rows: list[dict[str, Any]], headers: list[str], title: str = "Relatorio") -> None:
        try:
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            doc = SimpleDocTemplate(path, pagesize=landscape(A4))
            styles = getSampleStyleSheet()
            elements = [Paragraph(title, styles["Title"]), Spacer(1, 12)]

            table_data = [headers] + [[str(row.get(h, "")) for h in headers] for row in rows]
            table = Table(table_data, repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#d1d5db")),
                        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ]
                )
            )
            elements.append(table)
            doc.build(elements)
        except OSError as exc:
            raise ServiceError("Nao foi possivel salvar o arquivo PDF.", technical_detail=str(exc)) from exc
